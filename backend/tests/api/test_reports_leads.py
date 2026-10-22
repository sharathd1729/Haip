"""Reports: the rebuilt Leads report, agent access, and the download path.

The Reports section used to be a v1 export: leads filtered by a stage list the
CRM no longer has, no way to keep only leads who shared a phone number, every
download quietly capped at 500 rows, and owners/managers only. These tests pin
the new behaviour.
"""
import csv
import io
import uuid

from tests.conftest import seed_capabilities
from datetime import datetime, timezone

from app.core.security import create_access_token, get_password_hash
from app.models.audit_event import AuditEvent
from app.models.call import Call, CallDirection
from app.models.guest import Guest
from app.models.inquiry import Inquiry, InquirySource, InquiryStage
from app.models.pipeline_stage import PipelineStage
from app.models.tenant import Tenant
from app.models.user import User, UserRole

_HASH = get_password_hash("pw-for-tests")


def _tenant(db):
    t = Tenant(id=str(uuid.uuid4()), name="T", slug=f"rp-{uuid.uuid4().hex[:6]}", plan="growth")
    db.add(t)
    db.flush()
    seed_capabilities(db, t.id)       # the revenue report needs revenue.read, like a provisioned workspace has
    return t


def _user(db, tenant_id, role, name="Agent"):
    u = User(id=str(uuid.uuid4()), tenant_id=tenant_id, name=name,
             email=f"u-{uuid.uuid4().hex[:8]}@t.com", hashed_password=_HASH,
             role=role, is_active=True)
    db.add(u)
    db.flush()
    return u


def _headers(user):
    token = create_access_token({"sub": user.id, "tenant_id": user.tenant_id, "role": user.role.value})
    return {"Authorization": f"Bearer {token}"}


def _stage(db, tenant_id, key, label, order=0):
    s = PipelineStage(id=str(uuid.uuid4()), tenant_id=tenant_id, key=key, label=label, order=order)
    db.add(s)
    db.flush()
    return s


def _lead(db, tenant_id, *, name="Guest", phone="", email=None, source=InquirySource.META_IG,
          stage=None, owner=None, qualification="qualified", created_at=None, **extra):
    g = Guest(id=str(uuid.uuid4()), tenant_id=tenant_id, name=name, phone=phone, email=email)
    db.add(g)
    db.flush()
    inq = Inquiry(
        id=str(uuid.uuid4()), tenant_id=tenant_id, guest_id=g.id, stage=InquiryStage.NEW,
        source=source, stage_id=stage.id if stage else None,
        assigned_to=owner.id if owner else None, qualification_status=qualification,
        created_at=created_at or datetime.now(timezone.utc), **extra,
    )
    db.add(inq)
    db.flush()
    return inq


def _preview(client, user, fields, filters=None, **body):
    return client.post("/api/v1/reports/preview", headers=_headers(user), json={
        "report_id": "inquiries", "fields": fields, "filters": filters or {}, **body,
    })


# ── the filter the client asked for: leads who shared a number ───────────────

def test_only_leads_who_shared_a_phone_number(client, db):
    t = _tenant(db)
    owner = _user(db, t.id, UserRole.OWNER)
    _lead(db, t.id, name="Has Number", phone="+919876500001")
    _lead(db, t.id, name="Also Has", phone="+919876500002", source=InquirySource.INBOUND_CALL)
    _lead(db, t.id, name="Instagram Only", phone="")
    db.commit()

    res = _preview(client, owner, ["guest_name", "guest_phone"], {"has_phone": "true"})
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["total"] == 2
    assert {r["guest_name"] for r in body["rows"]} == {"Has Number", "Also Has"}

    no_number = _preview(client, owner, ["guest_name"], {"has_phone": "false"}).json()
    assert [r["guest_name"] for r in no_number["rows"]] == ["Instagram Only"]

    everyone = _preview(client, owner, ["guest_name"]).json()
    assert everyone["total"] == 3


def test_source_and_stage_filters_use_the_workspace_catalog(client, db):
    t = _tenant(db)
    owner = _user(db, t.id, UserRole.OWNER)
    active = _stage(db, t.id, "active", "Active", 1)
    rnr = _stage(db, t.id, "rnr", "RNR", 2)
    _lead(db, t.id, name="Call Active", source=InquirySource.INBOUND_CALL, stage=active)
    _lead(db, t.id, name="Call RNR", source=InquirySource.INBOUND_CALL, stage=rnr)
    _lead(db, t.id, name="DM Active", source=InquirySource.META_IG, stage=active)
    db.commit()

    res = _preview(client, owner, ["guest_name", "source", "stage"],
                   {"sources": ["inbound_call"], "stages": ["active"]}).json()
    assert [(r["guest_name"], r["source"], r["stage"]) for r in res["rows"]] == [
        ("Call Active", "Phone call (IVR)", "Active"),
    ]

    two = _preview(client, owner, ["guest_name"], {"sources": "inbound_call,meta_ig"}).json()
    assert two["total"] == 3


def test_unknown_source_is_rejected_not_ignored(client, db):
    t = _tenant(db)
    owner = _user(db, t.id, UserRole.OWNER)
    db.commit()
    res = _preview(client, owner, ["guest_name"], {"sources": ["carrier_pigeon"]})
    assert res.status_code == 422


def test_owner_filter_and_unassigned(client, db):
    t = _tenant(db)
    boss = _user(db, t.id, UserRole.OWNER)
    agent = _user(db, t.id, UserRole.CONCIERGE, "Asha")
    _lead(db, t.id, name="Asha's", owner=agent)
    _lead(db, t.id, name="Nobody's")
    db.commit()

    mine = _preview(client, boss, ["guest_name"], {"owner": agent.id}).json()
    assert [r["guest_name"] for r in mine["rows"]] == ["Asha's"]
    free = _preview(client, boss, ["guest_name"], {"owner": "unassigned"}).json()
    assert [r["guest_name"] for r in free["rows"]] == ["Nobody's"]


def test_spam_is_left_out_unless_asked_for(client, db):
    t = _tenant(db)
    owner = _user(db, t.id, UserRole.OWNER)
    _lead(db, t.id, name="Real")
    _lead(db, t.id, name="Spam", qualification="unqualified")
    db.commit()

    assert _preview(client, owner, ["guest_name"]).json()["total"] == 1
    assert _preview(client, owner, ["guest_name"], {"include_unqualified": "true"}).json()["total"] == 2


def test_dates_are_ist_days_inclusive(client, db):
    """31 Aug 19:00 UTC is 00:30 on 1 Sept in India; 30 Sept 19:00 UTC is
    already 1 Oct. September in IST is 1st 00:00 to 30th 23:59 IST."""
    t = _tenant(db)
    owner = _user(db, t.id, UserRole.OWNER)
    _lead(db, t.id, name="Sep 1 IST", created_at=datetime(2026, 8, 31, 19, 0, tzinfo=timezone.utc))
    _lead(db, t.id, name="Sep 30 IST", created_at=datetime(2026, 9, 30, 17, 0, tzinfo=timezone.utc))
    _lead(db, t.id, name="Oct 1 IST", created_at=datetime(2026, 9, 30, 19, 0, tzinfo=timezone.utc))
    _lead(db, t.id, name="Aug 31 IST", created_at=datetime(2026, 8, 31, 17, 0, tzinfo=timezone.utc))
    db.commit()

    res = _preview(client, owner, ["guest_name"], date_from="2026-09-01", date_to="2026-09-30").json()
    assert {r["guest_name"] for r in res["rows"]} == {"Sep 1 IST", "Sep 30 IST"}


# ── other workspaces stay out ────────────────────────────────────────────────

def test_another_workspaces_leads_never_appear(client, db):
    mine, theirs = _tenant(db), _tenant(db)
    owner = _user(db, mine.id, UserRole.OWNER)
    _lead(db, mine.id, name="Mine", phone="+919800000001")
    _lead(db, theirs.id, name="Theirs", phone="+919800000002")
    db.commit()

    res = _preview(client, owner, ["guest_name", "guest_phone"]).json()
    assert [r["guest_name"] for r in res["rows"]] == ["Mine"]
    assert res["total"] == 1


# ── agents: their own leads, nothing else ────────────────────────────────────

def test_an_agent_downloads_their_own_leads_and_unassigned_ones_like_the_leads_list(client, db):
    t = _tenant(db)
    asha = _user(db, t.id, UserRole.CONCIERGE, "Asha")
    ravi = _user(db, t.id, UserRole.RESERVATIONS, "Ravi")
    _lead(db, t.id, name="Asha lead", owner=asha, phone="+919811111111")
    _lead(db, t.id, name="Ravi lead", owner=ravi, phone="+919822222222")
    _lead(db, t.id, name="Unassigned", phone="+919833333333")
    db.commit()

    # One visibility rule (22 Oct 2026): the Leads list's "mine + unassigned", never a teammate's.
    res = _preview(client, asha, ["guest_name"]).json()
    assert sorted(r["guest_name"] for r in res["rows"]) == ["Asha lead", "Unassigned"]
    assert res["scope"] == "own"

    # Asking for someone else's leads doesn't widen it; asking for unassigned narrows it.
    sneaky = _preview(client, asha, ["guest_name"], {"owner": ravi.id}).json()
    assert sorted(r["guest_name"] for r in sneaky["rows"]) == ["Asha lead", "Unassigned"]
    only_free = _preview(client, asha, ["guest_name"], {"owner": "unassigned"}).json()
    assert [r["guest_name"] for r in only_free["rows"]] == ["Unassigned"]

    dl = client.post("/api/v1/reports/download?format=csv", headers=_headers(asha),
                     json={"report_id": "inquiries", "fields": ["guest_name", "guest_phone"]})
    assert dl.status_code == 200
    text = dl.text
    assert "Asha lead" in text and "Ravi lead" not in text and "919822222222" not in text


def test_an_agent_only_sees_their_own_calls(client, db):
    t = _tenant(db)
    asha = _user(db, t.id, UserRole.CONCIERGE, "Asha")
    ravi = _user(db, t.id, UserRole.RESERVATIONS, "Ravi")
    for who in (asha, ravi):
        db.add(Call(id=str(uuid.uuid4()), tenant_id=t.id, agent_id=who.id, ivr_provider="vobiz",
                    ivr_call_id=f"c-{uuid.uuid4().hex[:8]}", direction=CallDirection.INBOUND,
                    from_number="+919800000000", to_number="+918000000000", duration_seconds=30,
                    started_at="2026-09-10T10:00:00+00:00"))
    db.commit()

    res = client.post("/api/v1/reports/preview", headers=_headers(asha), json={
        "report_id": "calls", "fields": ["agent_name"],
        "date_from": "2026-09-01", "date_to": "2026-09-30",
    })
    assert res.status_code == 200, res.text
    assert [r["agent_name"] for r in res.json()["rows"]] == ["Asha"]


def test_agents_cannot_run_the_other_reports(client, db):
    t = _tenant(db)
    asha = _user(db, t.id, UserRole.CONCIERGE)
    db.commit()
    for report in ("revenue", "guests", "team", "qa_scores", "whatsapp"):
        res = client.post("/api/v1/reports/preview", headers=_headers(asha),
                          json={"report_id": report, "fields": ["date"]})
        assert res.status_code == 403, report


def test_other_roles_are_refused_and_see_an_empty_catalog(client, db):
    t = _tenant(db)
    for role in (UserRole.MARKETING, UserRole.READ_ONLY, UserRole.FOH):
        u = _user(db, t.id, role)
        db.commit()
        assert _preview(client, u, ["guest_name"]).status_code == 403
        assert client.post("/api/v1/reports/download?format=csv", headers=_headers(u),
                           json={"report_id": "inquiries", "fields": ["guest_name"]}).status_code == 403
        assert client.get("/api/v1/reports/catalog", headers=_headers(u)).json() == []


def test_catalog_is_role_scoped_and_carries_the_workspace_stages(client, db):
    t = _tenant(db)
    boss = _user(db, t.id, UserRole.OWNER, "Boss")
    asha = _user(db, t.id, UserRole.CONCIERGE, "Asha")
    _stage(db, t.id, "active", "Active", 1)
    _stage(db, t.id, "sold", "Sold", 2)
    db.commit()

    owner_cat = client.get("/api/v1/reports/catalog", headers=_headers(boss)).json()
    assert {r["id"] for r in owner_cat} >= {"inquiries", "calls", "revenue"}
    leads = next(r for r in owner_cat if r["id"] == "inquiries")
    flt = {f["key"]: f for f in leads["filters"]}
    assert "has_phone" in flt and flt["has_phone"]["type"] == "boolean"
    assert [o["value"] for o in flt["stages"]["options"]] == ["active", "sold"]
    assert {"unassigned", asha.id} <= {o["value"] for o in flt["owner"]["options"]}
    assert "inbound_call" in [o["value"] for o in flt["sources"]["options"]]

    agent_cat = client.get("/api/v1/reports/catalog", headers=_headers(asha)).json()
    assert {r["id"] for r in agent_cat} == {"inquiries", "calls"}
    agent_leads = next(r for r in agent_cat if r["id"] == "inquiries")
    assert "owner" not in {f["key"] for f in agent_leads["filters"]}


# ── the download itself ──────────────────────────────────────────────────────

def test_a_download_is_never_cut_at_500_rows_and_is_audited(client, db):
    t = _tenant(db)
    owner = _user(db, t.id, UserRole.OWNER)
    g = Guest(id=str(uuid.uuid4()), tenant_id=t.id, name="Bulk", phone="+919899999999")
    db.add(g)
    db.flush()
    db.add_all([
        Inquiry(id=str(uuid.uuid4()), tenant_id=t.id, guest_id=g.id, stage=InquiryStage.NEW,
                source=InquirySource.INBOUND_CALL, created_at=datetime.now(timezone.utc))
        for _ in range(520)
    ])
    db.commit()

    res = client.post("/api/v1/reports/download?format=csv", headers=_headers(owner), json={
        "report_id": "inquiries", "fields": ["guest_name", "guest_phone"],
        "filters": {"has_phone": "true"},
    })
    assert res.status_code == 200
    rows = list(csv.reader(io.StringIO(res.text)))
    assert rows[0] == ["Name", "Phone"]           # readable headers, not field keys
    assert len(rows) == 521

    event = db.query(AuditEvent).filter_by(tenant_id=t.id, action="reports.download").one()
    assert event.actor_user_id == owner.id
    assert event.payload["row_count"] == 520 and event.payload["scope"] == "workspace"
    assert event.payload["filters"] == {"has_phone": "true"}
    assert "919899999999" not in str(event.payload)     # the trail never holds the numbers


def test_excel_download_opens_and_has_the_columns_asked_for(client, db):
    from openpyxl import load_workbook

    t = _tenant(db)
    owner = _user(db, t.id, UserRole.OWNER)
    _lead(db, t.id, name="Priya", phone="+919876500003", source=InquirySource.INBOUND_CALL,
          ai_summary="Wants a caravan for 4 on the 12th.")
    db.commit()

    res = client.post("/api/v1/reports/download?format=xlsx", headers=_headers(owner), json={
        "report_id": "inquiries",
        "fields": ["guest_name", "guest_phone", "source", "ai_summary"],
    })
    assert res.status_code == 200
    ws = load_workbook(io.BytesIO(res.content)).active
    assert [c.value for c in ws[1]] == ["Name", "Phone", "Source", "Remarks (AI summary)"]
    assert [c.value for c in ws[2]] == ["Priya", "+919876500003", "Phone call (IVR)",
                                         "Wants a caravan for 4 on the 12th."]


def test_a_guest_typed_formula_cannot_run_in_the_spreadsheet(client, db):
    t = _tenant(db)
    owner = _user(db, t.id, UserRole.OWNER)
    _lead(db, t.id, name='=HYPERLINK("http://evil.example","click")', phone="+919876500004")
    _lead(db, t.id, name="@travelgram.priya", phone="")
    db.commit()

    res = client.post("/api/v1/reports/download?format=csv", headers=_headers(owner), json={
        "report_id": "inquiries", "fields": ["guest_name", "guest_phone"],
    })
    cells = {r[0]: r[1] for r in csv.reader(io.StringIO(res.text))}
    assert any(name.startswith("'=HYPERLINK") for name in cells)          # defused
    assert "@travelgram.priya" in cells                                    # a real handle survives
    assert cells['\'=HYPERLINK("http://evil.example","click")'] == "+919876500004"   # phone untouched


def test_unknown_report_and_empty_columns_are_clean_errors(client, db):
    t = _tenant(db)
    owner = _user(db, t.id, UserRole.OWNER)
    db.commit()
    h = _headers(owner)
    assert client.post("/api/v1/reports/preview", headers=h,
                       json={"report_id": "nope", "fields": ["x"]}).status_code == 404
    assert client.post("/api/v1/reports/preview", headers=h,
                       json={"report_id": "inquiries", "fields": ["not_a_column"]}).status_code == 422


def test_calls_has_transcript_filter_now_works(client, db):
    t = _tenant(db)
    owner = _user(db, t.id, UserRole.OWNER)
    for status in ("done", "pending"):
        db.add(Call(id=str(uuid.uuid4()), tenant_id=t.id, ivr_provider="vobiz",
                    ivr_call_id=f"c-{uuid.uuid4().hex[:8]}", direction=CallDirection.INBOUND,
                    from_number="+919800000000", to_number="+918000000000", duration_seconds=10,
                    transcription_status=status, started_at="2026-09-10T10:00:00+00:00"))
    db.commit()

    def run(value):
        return client.post("/api/v1/reports/preview", headers=_headers(owner), json={
            "report_id": "calls", "fields": ["transcription_status"],
            "date_from": "2026-09-01", "date_to": "2026-09-30", "filters": {"has_transcript": value},
        }).json()["rows"]

    assert [r["transcription_status"] for r in run("true")] == ["done"]
    assert [r["transcription_status"] for r in run("false")] == ["pending"]
