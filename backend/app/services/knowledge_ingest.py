"""Knowledge base bulk ingestion -- website crawling + document upload.

Both paths funnel into the same shared ingest_chunks(): extract text,
split into retrieval-sized pieces (knowledge_service.chunk_text, which
was already built for exactly this and never wired up), and store each
as a KnowledgeEntry with review_status="pending" -- a human must approve
before the AI can use any of it (knowledge_service.search() only
retrieves review_status="approved" rows), same safety gate manual entries
skip past because a person typed them directly.

Uploaded documents are stored in R2 via media_storage.py's existing
S3-compatible client (same one WhatsApp media/IVR audio/avatars already
use) -- never Supabase storage, which only ever holds the database in
this stack.
"""
from __future__ import annotations

import uuid
from dataclasses import dataclass, field
from datetime import datetime, timezone
from typing import Optional

import httpx
import structlog
from bs4 import BeautifulSoup
from sqlalchemy.orm import Session

from app.core.url_safety import UnsafeURLError, safe_get
from app.models.knowledge import KnowledgeEntry
from app.services import knowledge_service

log = structlog.get_logger()

# Generous enough for a real brochure/rate-card PDF, small enough that one
# bad upload can't tie up a request handler or balloon storage costs.
MAX_UPLOAD_MB = 15
MAX_UPLOAD_BYTES = MAX_UPLOAD_MB * 1024 * 1024
ALLOWED_UPLOAD_EXTS = {".pdf", ".docx", ".txt"}

# Website fetches happen inline in a request handler (like the Sarvam TTS
# call in tenants.py) -- bounded, fast enough not to need a Celery hop.
# A generous but finite cap stops a slow/hanging page from tying up the
# request indefinitely.
FETCH_TIMEOUT_SECONDS = 20.0
MAX_PAGE_BYTES = 5 * 1024 * 1024  # 5 MB -- a real content page, not a stream
MAX_URLS_PER_REQUEST = 5


class IngestError(Exception):
    """Raised for a real, user-facing ingestion failure (bad file, fetch
    failed, etc.) -- callers turn this into an HTTP 400/502 rather than a
    500, and can show the message directly."""


# ─── Text extraction ────────────────────────────────────────────────────────

def extract_text_from_pdf(content: bytes) -> str:
    from pypdf import PdfReader
    import io

    try:
        reader = PdfReader(io.BytesIO(content))
    except Exception as exc:
        raise IngestError(f"Couldn't read this PDF: {exc}") from exc
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            continue  # one bad page must not fail the whole document
    return "\n\n".join(p for p in pages if p.strip())


def extract_text_from_docx(content: bytes) -> str:
    import io
    from docx import Document

    try:
        doc = Document(io.BytesIO(content))
    except Exception as exc:
        raise IngestError(f"Couldn't read this Word document: {exc}") from exc
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def extract_text_from_html(html: str) -> str:
    """Strip nav/script/style/footer noise, keep the actual page content --
    a raw HTML dump would bury real sentences in menu links and JS."""
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header", "noscript", "svg"]):
        tag.decompose()
    text = soup.get_text(separator="\n")
    lines = [ln.strip() for ln in text.splitlines()]
    return "\n".join(ln for ln in lines if ln)


def extract_text_from_upload(filename: str, content_type: str, content: bytes) -> str:
    ext = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if ext == ".pdf" or "pdf" in content_type:
        return extract_text_from_pdf(content)
    if ext == ".docx" or "wordprocessingml" in content_type:
        return extract_text_from_docx(content)
    if ext == ".txt" or content_type.startswith("text/"):
        try:
            return content.decode("utf-8", errors="replace")
        except Exception as exc:
            raise IngestError(f"Couldn't read this file as text: {exc}") from exc
    raise IngestError(f"Unsupported file type '{ext or content_type}'. Use PDF, Word (.docx), or plain text.")


# ─── Website fetch ──────────────────────────────────────────────────────────

def fetch_url_text(url: str) -> str:
    if not url.startswith(("http://", "https://")):
        raise IngestError(f"'{url}' isn't a valid URL.")
    try:
        # safe_get refuses private/internal addresses on every redirect hop --
        # this URL is typed in by a tenant (SSRF).
        resp = safe_get(url, timeout=FETCH_TIMEOUT_SECONDS, headers={"User-Agent": "HAIP-KnowledgeBot/1.0"})
    except UnsafeURLError as exc:
        raise IngestError(f"Can't import {url}: {exc}") from exc
    except httpx.HTTPError as exc:
        raise IngestError(f"Couldn't reach {url}: {exc}") from exc
    if resp.status_code != 200:
        raise IngestError(f"{url} returned HTTP {resp.status_code}.")
    if len(resp.content) > MAX_PAGE_BYTES:
        raise IngestError(f"{url} is too large ({len(resp.content) / 1024 / 1024:.1f} MB) to ingest.")
    content_type = resp.headers.get("content-type", "")
    if "html" not in content_type and "<html" not in resp.text[:500].lower():
        raise IngestError(f"{url} doesn't look like an HTML page.")
    text = extract_text_from_html(resp.text)
    if not text.strip():
        raise IngestError(f"{url} had no readable text content.")
    return text


# ─── Shared ingest → pending KnowledgeEntry rows ───────────────────────────

@dataclass
class IngestReport:
    created: list[KnowledgeEntry] = field(default_factory=list)
    unchanged: int = 0          # passages already in the knowledge base from this same page/file; not added again
    marked_out_of_date: int = 0  # earlier passages from this page/file that the new version no longer contains


def _norm(text: str) -> str:
    return " ".join((text or "").split())


def ingest_chunks_report(
    db: Session,
    tenant_id: str,
    text: str,
    *,
    source: str,
    source_url: Optional[str] = None,
    source_ref: Optional[str] = None,
    file_s3_key: Optional[str] = None,
    category: str = "other",
    created_by: Optional[str] = None,
) -> IngestReport:
    """Chunk `text` and store each NEW piece as a pending KnowledgeEntry, queuing embedding for each.

    Never marks anything approved -- a human must review machine-sourced content before the AI can use it.

    Re-importing the same page or file is safe: a passage that is already there (same words) is not added a second time, and
    an earlier passage from that page/file that the new version no longer contains is marked out of date, so an old price or
    policy that was removed from the website stops being retrievable instead of living on beside the new text. Passages the
    owner wants to keep can be kept from the settings page.
    """
    chunks = knowledge_service.chunk_text(text)
    report = IngestReport()

    # The same origin = same tenant + source + page URL (website) or file name (document).
    origin = KnowledgeEntry.tenant_id == tenant_id
    q = db.query(KnowledgeEntry).filter(origin, KnowledgeEntry.source == source)
    if source_url:
        q = q.filter(KnowledgeEntry.source_url == source_url)
    elif source_ref:
        q = q.filter(KnowledgeEntry.source_ref == source_ref)
    else:
        q = q.filter(KnowledgeEntry.id == "")      # no stable origin: nothing to compare with
    earlier = q.all()
    by_text = {}
    for e in earlier:
        by_text.setdefault(_norm(e.content), e)

    seen_texts: set[str] = set()
    for chunk in chunks:
        key = _norm(chunk)
        if key in seen_texts:
            continue                                  # the page repeats itself
        seen_texts.add(key)
        match = by_text.get(key)
        if match is not None:
            report.unchanged += 1
            if match.stale_since is not None:
                match.stale_since = None              # it is back on the page
            continue
        entry = KnowledgeEntry(
            id=str(uuid.uuid4()),
            tenant_id=tenant_id,
            category=category,
            question=None,
            content=chunk,
            source=source,
            source_url=source_url,
            source_ref=source_ref,
            file_s3_key=file_s3_key,
            priority=0,
            is_active=True,
            review_status="pending",
            embedding_status="pending",
            created_by=created_by,
        )
        db.add(entry)
        report.created.append(entry)

    if chunks and earlier:
        now = datetime.now(timezone.utc)
        for e in earlier:
            if _norm(e.content) not in seen_texts and e.stale_since is None and e.review_status != "rejected":
                e.stale_since = now
                report.marked_out_of_date += 1

    db.commit()
    for entry in report.created:
        db.refresh(entry)
        try:
            from app.workers.tasks.embed_knowledge import embed_entry as embed_task
            embed_task.delay(tenant_id, entry.id)
        except Exception as exc:
            log.warning("kb_ingest_embed_enqueue_failed", entry_id=entry.id, tenant_id=tenant_id, error=str(exc))
    log.info(
        "kb_bulk_ingested", tenant_id=tenant_id, source=source, source_url=source_url,
        source_ref=source_ref, chunk_count=len(report.created), unchanged=report.unchanged,
        marked_out_of_date=report.marked_out_of_date,
    )
    return report


def ingest_chunks(db: Session, tenant_id: str, text: str, **kwargs) -> list[KnowledgeEntry]:
    """Back-compatible wrapper: the entries that were newly created."""
    return ingest_chunks_report(db, tenant_id, text, **kwargs).created
